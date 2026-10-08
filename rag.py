"""
Lightweight RAG trainer + test lab for "Level 1 - AI Fundamentals for Testers".

  python rag.py index        # read docs/ and build index.json
  python rag.py              # start the training chat

The app is deliberately transparent so trainees can TEST it: repeat runs
(/variance), compare models (/compare), inspect the pipeline (/trace),
break retrieval (/break), change temperature (/temp) and swap prompts (/prompt).

No vector DB, no LangChain: embeddings live in a JSON file and are searched
with plain cosine similarity (plenty for a curriculum).
"""
import json
import math
import random
import re
import sys
import time
import urllib.request
from itertools import combinations
from pathlib import Path

# ---- settings (tune here) -------------------------------------------------
OLLAMA_URL = "http://localhost:11434"
CHAT_MODEL = "llama3.2"            # 3B, ~2 GB RAM. Use "llama3.2:1b" for even less.
COMPARE_MODEL = "llama3.2:1b"      # second model for /compare (ollama pull llama3.2:1b)
EMBED_MODEL = "nomic-embed-text"   # 274 MB
DOCS_DIR = Path(__file__).parent / "docs"
INDEX_FILE = Path(__file__).parent / "index.json"
CHUNK_CHARS = 900                  # size of each text chunk
CHUNK_OVERLAP = 150
TOP_K = 4                          # chunks sent to the model per question
NUM_CTX = 4096                     # model context window; lower = less RAM
TRAINING_LEVEL = "Level 1 - AI Fundamentals for Testers"
FALLBACK = "This is not covered in the Level 1 material."

PROMPTS = {
    "strong": f"""You are a friendly, patient trainer delivering "{TRAINING_LEVEL}" to software
testers who are moving into testing AI systems. Teach ONLY from the curriculum context
provided. If the context does not contain the answer, reply exactly: "{FALLBACK}"
Use simple English, relate concepts to software testing, and give practical testing
examples. Always do exactly the TASK asked (a quiz must be a quiz, an exercise an exercise).""",
    # Deliberately weak prompt for the M1.3 exercise: no scope, grounding, fallback or format.
    "weak": "You are a helpful assistant. Answer the question.",
}

RETRIEVAL_MODES = {"normal": "normal retrieval", "off": "no context retrieved",
                   "noise": "random chunks retrieved", "worst": "least relevant chunks retrieved"}

LAB = {"temperature": 0.3, "retrieval": "normal", "prompt": "strong", "last_trace": None}


# ---- Ollama helpers -------------------------------------------------------
def ollama(path, payload=None, stream=False):
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(OLLAMA_URL + path, data=data,
                                 headers={"Content-Type": "application/json"})
    resp = urllib.request.urlopen(req, timeout=600)
    if not stream:
        return json.loads(resp.read())
    return (json.loads(line) for line in resp if line.strip())


def embed(texts):
    vectors = []
    for i in range(0, len(texts), 16):  # small batches keep memory low
        vectors += ollama("/api/embed", {"model": EMBED_MODEL, "input": texts[i:i + 16]})["embeddings"]
    return vectors


def installed_models():
    names = set()
    for m in ollama("/api/tags")["models"]:
        names.add(m["name"])
        names.add(m["name"].removesuffix(":latest"))
    return names


def generate(messages, model=CHAT_MODEL, echo=True, on_token=None):
    """Stream a chat completion; return (text, stats). on_token(token) is called per token (used by the UI)."""
    text, final, start = "", {}, time.time()
    for part in ollama("/api/chat", {
        "model": model, "messages": messages, "stream": True, "keep_alive": "10m",
        "options": {"num_ctx": NUM_CTX, "temperature": LAB["temperature"]},
    }, stream=True):
        token = part.get("message", {}).get("content", "")
        if echo:
            print(token, end="", flush=True)
        if on_token:
            on_token(token)
        text += token
        if part.get("done"):
            final = part
    return text, {"prompt_tokens": final.get("prompt_eval_count"),
                  "output_tokens": final.get("eval_count"),
                  "seconds": round(time.time() - start, 1)}


# ---- document loading -----------------------------------------------------
def read_docx(path):
    from docx import Document
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.lower().startswith(("heading", "title")):
            text = "## " + text
        lines.append(text)
    for table in doc.tables:  # curricula often live in tables
        for row in table.rows:
            cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_pdf(path):
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def read_file(path):
    ext = path.suffix.lower()
    if ext == ".docx":
        return read_docx(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def chunk(text):
    """Start a new chunk at each heading; pack paragraphs to ~CHUNK_CHARS; tag chunks with their heading."""
    chunks, current, heading = [], "", ""

    def flush():
        if current.strip():
            chunks.append((f"[{heading}]\n" if heading else "") + current.strip())

    for para in re.split(r"\n+", text):
        para = para.strip()
        if not para:
            continue
        if re.match(r"#{1,6}\s", para):
            flush()
            heading, current = para.lstrip("#").strip(), ""
            continue
        if len(current) + len(para) > CHUNK_CHARS and current:
            flush()
            current = current[-CHUNK_OVERLAP:]
        current += para + "\n"
    flush()
    return chunks


def build_index():
    files = sorted(p for p in DOCS_DIR.glob("*") if p.suffix.lower() in (".docx", ".pdf", ".txt", ".md"))
    if not files:
        sys.exit(f"No documents found. Put your material (.docx/.pdf/.txt/.md) in {DOCS_DIR}")
    records = []
    for f in files:
        pieces = chunk(read_file(f))
        print(f"  {f.name}: {len(pieces)} chunks")
        records += [{"source": f.name, "text": t} for t in pieces]
    print(f"Embedding {len(records)} chunks with {EMBED_MODEL}...")
    for rec, vec in zip(records, embed(["search_document: " + r["text"] for r in records])):
        rec["vec"] = [round(x, 5) for x in vec]
    INDEX_FILE.write_text(json.dumps(records), encoding="utf-8")
    print(f"Saved {INDEX_FILE.name}")


# ---- retrieval ------------------------------------------------------------
def cosine(a, b):
    dot = sum(x * y for x, y in zip(a, b))
    return dot / (math.sqrt(sum(x * x for x in a)) * math.sqrt(sum(y * y for y in b)) + 1e-9)


def retrieve(index, query, k=TOP_K):
    """Return [(score, record)]. LAB['retrieval'] can deliberately break this step."""
    mode = LAB["retrieval"]
    if mode == "off":
        return []
    if mode == "noise":
        return [(None, r) for r in random.sample(index, min(k, len(index)))]
    q = embed(["search_query: " + query])[0]
    scored = sorted(((cosine(q, r["vec"]), r) for r in index), key=lambda s: s[0], reverse=True)
    return scored[-k:][::-1] if mode == "worst" else scored[:k]


# ---- one RAG request ------------------------------------------------------
def ask(index, history, task, query, model=CHAT_MODEL, echo=True, k=TOP_K, hits=None, on_token=None):
    if hits is None:
        hits = retrieve(index, query, k)
    context = "\n\n---\n\n".join(r["text"] for _, r in hits) or "(no context retrieved)"
    if LAB["prompt"] == "strong":
        user_msg = f"Curriculum context:\n{context}\n\nTASK: {task}"
    else:
        user_msg = f"{context}\n\n{task}"
    messages = [{"role": "system", "content": PROMPTS[LAB["prompt"]]}] + (history or [])[-6:] + \
               [{"role": "user", "content": user_msg}]
    text, stats = generate(messages, model, echo, on_token)
    LAB["last_trace"] = {"model": model, "temperature": LAB["temperature"],
                         "retrieval": LAB["retrieval"], "prompt": LAB["prompt"],
                         "query": query, "hits": hits, "messages": messages,
                         "answer": text, **stats}
    if history is not None:
        history += [{"role": "user", "content": task}, {"role": "assistant", "content": text}]
    return text, hits, stats


def training_task(index, cmd, arg, text):
    """Turn a training command (or plain question) into (task, query, k, fixed_hits) for ask()."""
    if cmd == "/topics":
        # an outline needs EVERY module, which top-k similarity can't guarantee
        outline = [(1.0, r) for r in index if re.match(r"\[(module|level 1 labs)", r["text"], re.I)]
        fixed_hits = outline if LAB["retrieval"] == "normal" and outline else None
        task = ("List the Level 1 modules (M1.1 to M1.4) with their hours and main topics "
                "as a numbered outline, then list the Level 1 labs.")
        return task, "Level 1 modules M1.1 M1.2 M1.3 M1.4 topics labs", 6, fixed_hits
    if cmd == "/teach":
        task = (f"Teach '{arg}' as a short lesson for testers: objective, key concepts, "
                "a simple example, why it matters for testing AI systems, and a summary. "
                "End with one check-your-understanding question.")
    elif cmd == "/quiz":
        task = (f"Write a 5-question multiple-choice quiz on '{arg}' for testers. Use "
                "scenario-based questions where possible. List the answers with a one-line "
                "explanation at the end.")
    elif cmd == "/exercise":
        task = (f"Give one hands-on exercise on '{arg}' that a tester can do, with steps, "
                "what to record, and the expected learning outcome.")
    else:
        task = (f"Answer the trainee's question with a practical testing example, then ask "
                f"one check-your-understanding question.\nQuestion: {text}")
        return task, text, TOP_K, None
    return task, arg, TOP_K, None


def print_footer(hits, stats):
    sources = sorted({r["source"] for _, r in hits}) or ["none"]
    print(f"\n\n  (sources: {', '.join(sources)} | {stats['output_tokens']} tokens, "
          f"{stats['seconds']}s | /trace for details)\n")


# ---- lab commands ---------------------------------------------------------
def words(text):
    return set(re.findall(r"[a-z0-9']+", text.lower()))


def variance_report(answers):
    """Compare repeated answers to the same question: exact, semantic and word-level agreement."""
    vecs = embed(["clustering: " + a for a in answers])
    pairs = list(combinations(range(len(answers)), 2))
    semantic = sum(cosine(vecs[i], vecs[j]) for i, j in pairs) / len(pairs)
    overlap = sum(len(words(answers[i]) & words(answers[j])) / max(1, len(words(answers[i]) | words(answers[j])))
                  for i, j in pairs) / len(pairs)
    return {"distinct": len(set(answers)), "semantic": semantic, "overlap": overlap,
            "lengths": [len(a.split()) for a in answers]}


def cmd_variance(index, arg):
    m = re.match(r"(\d+)\s+(.*)", arg)
    n, question = (int(m.group(1)), m.group(2)) if m else (5, arg)
    if not question:
        print("Usage: /variance [runs] <question>   e.g. /variance 5 What is a token?\n")
        return
    n = max(2, min(n, 10))
    print(f"\nRunning the same question {n} times (temperature {LAB['temperature']}, "
          f"prompt {LAB['prompt']}, retrieval {LAB['retrieval']})...\n")
    answers, times = [], []
    for i in range(n):
        text, _, stats = ask(index, None, question, question, echo=False)
        answers.append(text.strip())
        times.append(stats["seconds"])
        preview = " ".join(text.split())[:220]
        print(f"  Run {i + 1} ({stats['output_tokens']} tokens, {stats['seconds']}s): {preview}...\n")
    v = variance_report(answers)
    lengths = v["lengths"]
    print("  ---- Variance report ----")
    print(f"  Distinct answers (exact text): {v['distinct']} of {n}")
    print(f"  Avg semantic similarity:       {v['semantic']:.3f}   (1.0 = same meaning)")
    print(f"  Avg word overlap (Jaccard):    {v['overlap']:.3f}   (1.0 = same words)")
    print(f"  Answer length (words):         min {min(lengths)}, max {max(lengths)}")
    print(f"  Latency (s):                   min {min(times)}, max {max(times)}")
    print("  Ask yourself: would an exact-match assertion pass? What threshold would you set?\n"
          "  Try again after /temp 0 or /prompt weak and compare.\n")


def cmd_compare(index, question):
    if not question:
        print("Usage: /compare <question>\n")
        return
    available = installed_models()
    models = [CHAT_MODEL, COMPARE_MODEL]
    missing = [m for m in models if m not in available]
    if missing:
        print(f"Model(s) not installed: {', '.join(missing)}. Run: ollama pull {missing[0]}\n"
              f"(or change COMPARE_MODEL in rag.py to one of: {', '.join(sorted(m for m in available if 'embed' not in m))})\n")
        return
    for model in models:
        print(f"\n===== {model} =====\n")
        _, hits, stats = ask(index, None, question, question, model=model)
        print(f"\n\n  ({stats['output_tokens']} output tokens, {stats['seconds']}s)")
    print("\nDocument: correctness, constraint-following, length, tone, tokens and latency.\n")


def cmd_trace(arg):
    t = LAB["last_trace"]
    if not t:
        print("Nothing to trace yet. Ask a question first.\n")
        return
    print("\n---- Trace of the last request ----")
    print(f"  1. Settings   model={t['model']}  temperature={t['temperature']}  "
          f"prompt={t['prompt']}  retrieval={t['retrieval']}")
    print(f"  2. Query embedded with {EMBED_MODEL}: {t['query'][:100]!r}")
    print(f"  3. Retrieved {len(t['hits'])} chunk(s) from {INDEX_FILE.name}:")
    for score, r in t["hits"]:
        s = "random" if score is None else f"{score:.3f}"
        first = " ".join(r["text"].split())[:110]
        print(f"     [{s}] {r['source']}: {first}...")
    print(f"  4. Prompt sent: {len(t['messages'])} messages, {t['prompt_tokens']} prompt tokens")
    print(f"  5. Output: {t['output_tokens']} tokens in {t['seconds']}s")
    if arg.strip() == "full":
        for m in t["messages"]:
            print(f"\n  [{m['role'].upper()}]\n{m['content']}")
    else:
        print("  (use /trace full to print the exact prompt sent to the model)")
    print()


def cmd_settings(cmd, arg):
    if cmd == "/temp":
        try:
            LAB["temperature"] = max(0.0, min(2.0, float(arg)))
            print(f"Temperature set to {LAB['temperature']}\n")
        except ValueError:
            print(f"Usage: /temp 0.0 - 2.0   (current {LAB['temperature']})\n")
    elif cmd == "/break":
        if arg not in ("off", "noise", "worst", "fix"):
            print("Usage: /break off | noise | worst | fix\n")
            return
        LAB["retrieval"] = "normal" if arg == "fix" else arg
        print(f"Retrieval mode: {LAB['retrieval']} ({RETRIEVAL_MODES[LAB['retrieval']]})\n")
    elif cmd == "/prompt":
        if arg in PROMPTS:
            LAB["prompt"] = arg
            print(f"System prompt switched to '{arg}'\n")
        elif arg == "show":
            print(f"\n[{LAB['prompt']} system prompt]\n{PROMPTS[LAB['prompt']]}\n")
        else:
            print("Usage: /prompt strong | weak | show\n")


HELP = """Training:
  /topics              list the Level 1 modules and topics
  /teach <topic>       short lesson on a topic          e.g. /teach temperature
  /quiz <topic>        5 multiple-choice questions       e.g. /quiz RAG
  /exercise <topic>    hands-on testing exercise
  <anything else>      ask a question
Test lab (treat this app as the AI system under test):
  /variance [n] <q>    run the same question n times and measure variance   (M1.1, Lab 2)
  /compare <q>         same question on two models                          (M1.1, Lab 2)
  /trace [full]        show retrieved chunks, scores, tokens, latency        (M1.2, Lab 1)
  /break off|noise|worst|fix   deliberately break retrieval                  (M1.2)
  /temp <0-2>          change temperature                                    (M1.1)
  /prompt strong|weak|show     swap the system prompt                        (M1.3)
  /status  /reset  /help  /quit"""


def chat():
    if not INDEX_FILE.exists():
        sys.exit("No index yet. Run:  python rag.py index")
    index = json.loads(INDEX_FILE.read_text(encoding="utf-8"))
    history = []
    print(f"{TRAINING_LEVEL} trainer ready ({len(index)} chunks, model {CHAT_MODEL}).\n\n{HELP}\n")
    while True:
        try:
            text = input("You> ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if not text:
            continue
        cmd, _, arg = text.partition(" ")
        cmd, arg = cmd.lower(), arg.strip()
        try:
            if cmd in ("/quit", "/exit"):
                break
            elif cmd == "/help":
                print(HELP + "\n")
            elif cmd == "/reset":
                history.clear()
                print("Conversation cleared.\n")
            elif cmd == "/status":
                print(f"model={CHAT_MODEL} temperature={LAB['temperature']} prompt={LAB['prompt']} "
                      f"retrieval={LAB['retrieval']} history={len(history) // 2} turns\n")
            elif cmd in ("/temp", "/break", "/prompt"):
                cmd_settings(cmd, arg.lower())
            elif cmd == "/trace":
                cmd_trace(arg)
            elif cmd == "/variance":
                cmd_variance(index, arg)
            elif cmd == "/compare":
                cmd_compare(index, arg)
            elif cmd.startswith("/") and cmd not in ("/topics", "/teach", "/quiz", "/exercise"):
                print("Unknown command. Type /help\n")
            else:
                task, query, k, fixed_hits = training_task(index, cmd, arg, text)
                print("\nTrainer> ", end="")
                _, hits, stats = ask(index, history, task, query, k=k, hits=fixed_hits)
                print_footer(hits, stats)
        except Exception as e:
            hint = "Is Ollama running? Try: ollama serve" if isinstance(e, OSError) else ""
            print(f"\n[error] {e}\n{hint}\n")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) > 1 and sys.argv[1] == "index":
        build_index()
    else:
        chat()
